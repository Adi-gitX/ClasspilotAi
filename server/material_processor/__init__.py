"""
Material Processor - PDF, PPT, Image processing for ClassPilot AI
Extracts text content for RAG context
"""

import os
import io
import uuid
from typing import Optional, Dict, List, Tuple
from pathlib import Path

# PDF processing
try:
    from PyPDF2 import PdfReader
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

# PPT processing
try:
    from pptx import Presentation
    HAS_PPT = True
except ImportError:
    HAS_PPT = False

# OCR for images
try:
    from PIL import Image
    import pytesseract
    HAS_OCR = True
except ImportError:
    HAS_OCR = False


class MaterialProcessor:
    """Processes uploaded materials (PDF, PPT, images) to extract text for RAG"""
    
    SUPPORTED_TYPES = {
        "pdf": [".pdf"],
        "ppt": [".ppt", ".pptx"],
        "image": [".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"],
        "text": [".txt", ".md", ".markdown"]
    }
    
    def __init__(self, upload_dir: str = "./data/uploads"):
        self.upload_dir = Path(upload_dir)
        self.upload_dir.mkdir(parents=True, exist_ok=True)
    
    def get_file_type(self, filename: str) -> Optional[str]:
        """Determine file type from extension"""
        ext = Path(filename).suffix.lower()
        for file_type, extensions in self.SUPPORTED_TYPES.items():
            if ext in extensions:
                return file_type
        return None
    
    def process(self, file_bytes: bytes, filename: str, save: bool = True) -> Dict:
        """
        Process uploaded file and extract text content
        
        Returns:
            {
                "id": str,
                "filename": str,
                "type": str,
                "content": str,
                "size": int,
                "path": str (if saved),
                "success": bool,
                "error": str (if failed)
            }
        """
        file_id = str(uuid.uuid4())
        file_type = self.get_file_type(filename)
        
        result = {
            "id": file_id,
            "filename": filename,
            "type": file_type or "unknown",
            "size": len(file_bytes),
            "success": False,
            "content": "",
            "error": None
        }
        
        if not file_type:
            result["error"] = f"Unsupported file type: {Path(filename).suffix}"
            return result
        
        # Save file if requested
        if save:
            file_path = self.upload_dir / f"{file_id}_{filename}"
            file_path.write_bytes(file_bytes)
            result["path"] = str(file_path)
        
        # Extract content based on type
        try:
            if file_type == "pdf":
                result["content"] = self._process_pdf(file_bytes)
            elif file_type == "ppt":
                result["content"] = self._process_ppt(file_bytes)
            elif file_type == "image":
                result["content"] = self._process_image(file_bytes)
            elif file_type == "text":
                result["content"] = file_bytes.decode("utf-8", errors="ignore")
            
            result["success"] = True
            
        except Exception as e:
            result["error"] = str(e)
        
        return result
    
    def _process_pdf(self, file_bytes: bytes) -> str:
        """Extract text from PDF"""
        if not HAS_PDF:
            return "[PDF processing not available - install PyPDF2]"
        
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            text_parts = []
            
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(f"## Page {i + 1}\n{page_text}")
            
            return "\n\n".join(text_parts) if text_parts else "[No text found in PDF]"
            
        except Exception as e:
            return f"[PDF processing error: {e}]"
    
    def _process_ppt(self, file_bytes: bytes) -> str:
        """Extract text from PowerPoint"""
        if not HAS_PPT:
            return "[PPT processing not available - install python-pptx]"
        
        try:
            prs = Presentation(io.BytesIO(file_bytes))
            text_parts = []
            
            for i, slide in enumerate(prs.slides):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_texts.append(shape.text)
                
                if slide_texts:
                    text_parts.append(f"## Slide {i + 1}\n" + "\n".join(slide_texts))
            
            return "\n\n".join(text_parts) if text_parts else "[No text found in PPT]"
            
        except Exception as e:
            return f"[PPT processing error: {e}]"
    
    def _process_image(self, file_bytes: bytes) -> str:
        """Extract text from image using OCR"""
        if not HAS_OCR:
            return "[OCR not available - install pytesseract and PIL]"
        
        try:
            image = Image.open(io.BytesIO(file_bytes))
            text = pytesseract.image_to_string(image)
            return text.strip() if text.strip() else "[No text found in image]"
            
        except Exception as e:
            return f"[OCR error: {e}]"
    
    def get_content_preview(self, content: str, max_length: int = 500) -> str:
        """Get a preview of the content"""
        if len(content) <= max_length:
            return content
        return content[:max_length] + "..."


# Singleton instance
processor = MaterialProcessor()


def process_material(file_bytes: bytes, filename: str) -> Dict:
    """Convenience function to process a material"""
    return processor.process(file_bytes, filename)


__all__ = ["MaterialProcessor", "processor", "process_material"]
